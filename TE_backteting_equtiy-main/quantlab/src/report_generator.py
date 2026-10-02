"""Comprehensive Multi-Sheet Master Backtest Report Generator."""

from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from .config import BacktestConfig
from .data_loader import ValidationReport
from .drawdown import DrawdownMetrics
from .metrics import BacktestMetrics


def generate_master_backtest_report(
    trades: pd.DataFrame,
    daily_df: pd.DataFrame,
    monthly_df: pd.DataFrame,
    yearly_df: pd.DataFrame,
    metrics: BacktestMetrics,
    val_report: ValidationReport,
    reconciliation_df: pd.DataFrame,
    llm_audit: Dict[str, Any],
    llm_text: str,
    config: BacktestConfig,
    output_path: str
) -> None:
    """
    Generates output/backtest_report.xlsx with 10 required sheets:
    1. Summary
    2. Trade Log
    3. Monthly Returns
    4. Yearly Returns
    5. Equity Curve
    6. Drawdown
    7. Metrics
    8. Input Validation
    9. Excel vs Python Reconciliation
    10. LLM Validation
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    # Standard styling tokens
    primary_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    accent_green = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    accent_red = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_code = Font(name="Courier New", size=9, color="374151")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    def auto_fit_columns(ws, min_w=12, max_w=50):
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    val = max(val.split("\n"), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_w), max_w)

    # =========================================================================
    # 1. SUMMARY SHEET
    # =========================================================================
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum.merge_cells("A1:D1")
    ws_sum["A1"] = "QUANTLAB BACKTEST SUMMARY REPORT"
    ws_sum["A1"].fill = primary_fill
    ws_sum["A1"].font = font_title
    ws_sum["A1"].alignment = align_center

    summary_items = [
        ("Initial Capital", f"₹{metrics.initial_capital:,.2f}"),
        ("Final Equity", f"₹{metrics.final_equity:,.2f}"),
        ("Total P&L", f"₹{metrics.total_pnl:,.2f}"),
        ("Total Return", f"{metrics.total_return_pct:.2f}%"),
        ("CAGR", f"{metrics.cagr * 100:.2f}%"),
        ("Total Trades", f"{metrics.total_trades}"),
        ("Win Rate", f"{metrics.win_rate * 100:.2f}%"),
        ("Profit Factor", f"{metrics.profit_factor:.2f}"),
        ("Expectancy", f"₹{metrics.expectancy:,.2f}"),
        ("Sharpe", f"{metrics.sharpe_ratio:.2f}"),
        ("Sortino", f"{metrics.sortino_ratio:.2f}"),
        ("Calmar", f"{metrics.calmar_ratio:.2f}"),
        ("Maximum Drawdown", f"₹{metrics.max_drawdown:,.2f}"),
        ("Maximum Drawdown %", f"{metrics.max_drawdown_pct:.2f}%"),
        ("Maximum Drawdown Duration", f"{metrics.max_drawdown_duration} days"),
        ("Skewness", f"{metrics.skewness:.4f}"),
        ("Max Win Streak", f"{metrics.max_win_streak}"),
        ("Max Loss Streak", f"{metrics.max_loss_streak}")
    ]

    for idx, (label, val_str) in enumerate(summary_items, 3):
        c1 = ws_sum.cell(row=idx, column=2, value=label)
        c1.font = font_bold
        c1.border = border_thin
        c1.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        c2 = ws_sum.cell(row=idx, column=3, value=val_str)
        c2.font = font_bold
        c2.border = border_thin
        c2.alignment = align_right

    ws_sum.column_dimensions["A"].width = 6
    ws_sum.column_dimensions["B"].width = 30
    ws_sum.column_dimensions["C"].width = 25
    ws_sum.column_dimensions["D"].width = 6

    # =========================================================================
    # 2. TRADE LOG SHEET
    # =========================================================================
    ws_log = wb.create_sheet(title="Trade Log")
    ws_log.views.sheetView[0].showGridLines = True

    log_headers = [
        "Trade ID", "Stock", "Entry Date", "Exit Date", "Duration", "Direction",
        "Entry Price", "Exit Price", "Quantity", "Allocated Capital", "Gross P&L",
        "Return %", "Cumulative P&L", "Cumulative Return %", "Winning Trade", "Losing Trade"
    ]
    for c_idx, h in enumerate(log_headers, 1):
        cell = ws_log.cell(row=1, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    for r_idx, (_, row) in enumerate(trades.iterrows(), 2):
        ws_log.cell(row=r_idx, column=1, value=row["trade_id"]).border = border_thin
        ws_log.cell(row=r_idx, column=2, value=row["stock"]).border = border_thin
        ws_log.cell(row=r_idx, column=3, value=row["entry_date"].strftime("%Y-%m-%d")).border = border_thin
        ws_log.cell(row=r_idx, column=4, value=row["exit_date"].strftime("%Y-%m-%d")).border = border_thin
        ws_log.cell(row=r_idx, column=5, value=int(row["duration"])).border = border_thin
        ws_log.cell(row=r_idx, column=6, value=row["direction"]).border = border_thin
        ws_log.cell(row=r_idx, column=7, value=round(float(row["entry_price"]), 2)).border = border_thin
        ws_log.cell(row=r_idx, column=8, value=round(float(row["exit_price"]), 2)).border = border_thin
        ws_log.cell(row=r_idx, column=9, value=round(float(row["quantity"]), 4)).border = border_thin
        ws_log.cell(row=r_idx, column=10, value=round(float(row["allocated_capital"]), 2)).border = border_thin
        ws_log.cell(row=r_idx, column=11, value=round(float(row["gross_pnl"]), 2)).border = border_thin
        ws_log.cell(row=r_idx, column=12, value=round(float(row["calc_return_pct"]), 2)).border = border_thin
        ws_log.cell(row=r_idx, column=13, value=round(float(row["cumulative_pnl"]), 2)).border = border_thin
        ws_log.cell(row=r_idx, column=14, value=round(float(row["cumulative_return_pct"]), 2)).border = border_thin
        ws_log.cell(row=r_idx, column=15, value=bool(row["winning_trade"])).border = border_thin
        ws_log.cell(row=r_idx, column=16, value=bool(row["losing_trade"])).border = border_thin

    auto_fit_columns(ws_log)

    # =========================================================================
    # 3. MONTHLY RETURNS SHEET
    # =========================================================================
    ws_mo = wb.create_sheet(title="Monthly Returns")
    ws_mo.views.sheetView[0].showGridLines = True

    for c_idx, col_name in enumerate(monthly_df.columns, 1):
        cell = ws_mo.cell(row=1, column=c_idx, value=col_name)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    for r_idx, (_, row) in enumerate(monthly_df.iterrows(), 2):
        for c_idx, col_name in enumerate(monthly_df.columns, 1):
            val = row[col_name]
            cell = ws_mo.cell(row=r_idx, column=c_idx, value=val if pd.notna(val) else "")
            cell.border = border_thin
            cell.alignment = align_right if c_idx > 1 else align_center
            if c_idx > 1 and pd.notna(val):
                if val > 0:
                    cell.fill = accent_green
                elif val < 0:
                    cell.fill = accent_red

    auto_fit_columns(ws_mo)

    # =========================================================================
    # 4. YEARLY RETURNS SHEET
    # =========================================================================
    ws_yr = wb.create_sheet(title="Yearly Returns")
    ws_yr.views.sheetView[0].showGridLines = True

    for c_idx, col_name in enumerate(yearly_df.columns, 1):
        cell = ws_yr.cell(row=1, column=c_idx, value=col_name)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    for r_idx, (_, row) in enumerate(yearly_df.iterrows(), 2):
        for c_idx, col_name in enumerate(yearly_df.columns, 1):
            val = row[col_name]
            cell = ws_yr.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border_thin
            cell.alignment = align_right if c_idx > 1 else align_center

    auto_fit_columns(ws_yr)

    # =========================================================================
    # 5. EQUITY CURVE SHEET
    # =========================================================================
    ws_eq = wb.create_sheet(title="Equity Curve")
    ws_eq.views.sheetView[0].showGridLines = True

    eq_cols = ["Date", "Equity", "Cumulative P&L", "Cumulative Return %"]
    for c_idx, col_name in enumerate(eq_cols, 1):
        cell = ws_eq.cell(row=1, column=c_idx, value=col_name)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    for r_idx, (_, row) in enumerate(daily_df.iterrows(), 2):
        ws_eq.cell(row=r_idx, column=1, value=row["date"].strftime("%Y-%m-%d")).border = border_thin
        ws_eq.cell(row=r_idx, column=2, value=round(float(row["equity"]), 2)).border = border_thin
        ws_eq.cell(row=r_idx, column=3, value=round(float(row["cumulative_pnl"]), 2)).border = border_thin
        ws_eq.cell(row=r_idx, column=4, value=round(float(row["cumulative_return_pct"]), 2)).border = border_thin

    auto_fit_columns(ws_eq)

    # =========================================================================
    # 6. DRAWDOWN SHEET
    # =========================================================================
    ws_dd = wb.create_sheet(title="Drawdown")
    ws_dd.views.sheetView[0].showGridLines = True

    dd_cols = ["Date", "Equity", "Running Peak", "Drawdown (₹)", "Drawdown %"]
    for c_idx, col_name in enumerate(dd_cols, 1):
        cell = ws_dd.cell(row=1, column=c_idx, value=col_name)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    for r_idx, (_, row) in enumerate(daily_df.iterrows(), 2):
        ws_dd.cell(row=r_idx, column=1, value=row["date"].strftime("%Y-%m-%d")).border = border_thin
        ws_dd.cell(row=r_idx, column=2, value=round(float(row["equity"]), 2)).border = border_thin
        ws_dd.cell(row=r_idx, column=3, value=round(float(row["running_peak"]), 2)).border = border_thin
        ws_dd.cell(row=r_idx, column=4, value=round(float(row["drawdown"]), 2)).border = border_thin
        ws_dd.cell(row=r_idx, column=5, value=round(float(row["drawdown_pct"]), 2)).border = border_thin

    auto_fit_columns(ws_dd)

    # =========================================================================
    # 7. METRICS SHEET
    # =========================================================================
    ws_m = wb.create_sheet(title="Metrics")
    ws_m.views.sheetView[0].showGridLines = True

    m_headers = ["Category", "Metric", "Value"]
    for c_idx, h in enumerate(m_headers, 1):
        cell = ws_m.cell(row=1, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    detailed_metrics = [
        # Trade Statistics
        ("Trade Statistics", "Total Trades", metrics.total_trades),
        ("Trade Statistics", "Winning Trades", metrics.winning_trades),
        ("Trade Statistics", "Losing Trades", metrics.losing_trades),
        ("Trade Statistics", "Even Trades", metrics.even_trades),
        ("Trade Statistics", "Win Rate", f"{metrics.win_rate*100:.2f}%"),
        ("Trade Statistics", "Loss Rate", f"{metrics.loss_rate*100:.2f}%"),

        # Returns
        ("Returns", "Initial Capital", f"₹{metrics.initial_capital:,.2f}"),
        ("Returns", "Final Equity", f"₹{metrics.final_equity:,.2f}"),
        ("Returns", "Total P&L", f"₹{metrics.total_pnl:,.2f}"),
        ("Returns", "Realized P&L", f"₹{metrics.realized_pnl:,.2f}"),
        ("Returns", "Total Return %", f"{metrics.total_return_pct:.2f}%"),
        ("Returns", "Average P&L per Trade", f"₹{metrics.avg_pnl_per_trade:,.2f}"),
        ("Returns", "Median P&L", f"₹{metrics.median_pnl:,.2f}"),
        ("Returns", "CAGR", f"{metrics.cagr*100:.2f}%"),
        ("Returns", "Trading Years", f"{metrics.years:.2f} yrs"),

        # Profitability
        ("Profitability", "Gross Profit", f"₹{metrics.gross_profit:,.2f}"),
        ("Profitability", "Gross Loss", f"₹{metrics.gross_loss:,.2f}"),
        ("Profitability", "Profit Factor", f"{metrics.profit_factor:.2f}"),
        ("Profitability", "Average Win", f"₹{metrics.avg_win:,.2f}"),
        ("Profitability", "Average Loss", f"₹{metrics.avg_loss:,.2f}"),
        ("Profitability", "Risk/Reward Ratio", f"{metrics.risk_reward_ratio:.2f}"),
        ("Profitability", "Expectancy", f"₹{metrics.expectancy:,.2f}"),

        # Risk Metrics
        ("Risk Metrics", "Sharpe Ratio", f"{metrics.sharpe_ratio:.2f}"),
        ("Risk Metrics", "Sortino Ratio", f"{metrics.sortino_ratio:.2f}"),
        ("Risk Metrics", "Calmar Ratio", f"{metrics.calmar_ratio:.2f}"),
        ("Risk Metrics", "Maximum Drawdown", f"₹{metrics.max_drawdown:,.2f}"),
        ("Risk Metrics", "Maximum Drawdown %", f"{metrics.max_drawdown_pct:.2f}%"),
        ("Risk Metrics", "Max Drawdown Duration", f"{metrics.max_drawdown_duration} days"),
        ("Risk Metrics", "Drawdown Peak Date", metrics.drawdown_peak_date),
        ("Risk Metrics", "Drawdown Trough Date", metrics.drawdown_trough_date),
        ("Risk Metrics", "Recovery Date", metrics.recovery_date),
        ("Risk Metrics", "Volatility (Ann. Daily)", f"{metrics.volatility*100:.2f}%"),

        # Skewness & Streaks
        ("Distribution & Streaks", "Skewness", f"{metrics.skewness:.4f}"),
        ("Distribution & Streaks", "Mean P&L", f"₹{metrics.mean_pnl:,.2f}"),
        ("Distribution & Streaks", "P&L Std Dev", f"₹{metrics.pnl_std:,.2f}"),
        ("Distribution & Streaks", "Max Win Streak", metrics.max_win_streak),
        ("Distribution & Streaks", "Max Loss Streak", metrics.max_loss_streak),
        ("Distribution & Streaks", "Max Consec Profit", f"₹{metrics.max_consecutive_profit:,.2f}"),
        ("Distribution & Streaks", "Max Consec Loss", f"₹{metrics.max_consecutive_loss:,.2f}")
    ]

    for r_idx, (cat, name, val) in enumerate(detailed_metrics, 2):
        ws_m.cell(row=r_idx, column=1, value=cat).border = border_thin
        ws_m.cell(row=r_idx, column=2, value=name).border = border_thin
        ws_m.cell(row=r_idx, column=3, value=str(val)).border = border_thin

    auto_fit_columns(ws_m)

    # =========================================================================
    # 8. INPUT VALIDATION SHEET
    # =========================================================================
    ws_iv = wb.create_sheet(title="Input Validation")
    ws_iv.views.sheetView[0].showGridLines = True

    iv_headers = ["Validation Item", "Result / Value"]
    for c_idx, h in enumerate(iv_headers, 1):
        cell = ws_iv.cell(row=1, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    iv_items = [
        ("Input File Path", val_report.file_path),
        ("Active Sheet", val_report.sheet_name),
        ("Total Raw Rows", val_report.total_raw_rows),
        ("Parsed Valid Trades", val_report.data_rows),
        ("Detected Columns", ", ".join(val_report.columns)),
        ("Unique Stocks Count", val_report.unique_stocks),
        ("Start Date", val_report.start_date or "N/A"),
        ("End Date", val_report.end_date or "N/A"),
        ("Active / Unclosed Trades", val_report.active_trades),
        ("Missing Values Count", val_report.missing_values),
        ("Duplicate Trades Count", val_report.duplicate_trades),
        ("Invalid Prices Count", val_report.invalid_prices),
        ("Invalid Dates Count", val_report.invalid_dates),
        ("Invalid Returns Count", val_report.invalid_returns)
    ]
    for r_idx, (k, v) in enumerate(iv_items, 2):
        ws_iv.cell(row=r_idx, column=1, value=k).border = border_thin
        ws_iv.cell(row=r_idx, column=2, value=str(v)).border = border_thin

    auto_fit_columns(ws_iv)

    # =========================================================================
    # 9. EXCEL VS PYTHON RECONCILIATION SHEET
    # =========================================================================
    ws_rec = wb.create_sheet(title="Excel vs Python Reconciliation")
    ws_rec.views.sheetView[0].showGridLines = True

    for c_idx, col_name in enumerate(reconciliation_df.columns, 1):
        cell = ws_rec.cell(row=1, column=c_idx, value=col_name)
        cell.fill = header_fill
        cell.font = font_header
        cell.alignment = align_center
        cell.border = border_thin

    for r_idx, (_, row) in enumerate(reconciliation_df.iterrows(), 2):
        for c_idx, col_name in enumerate(reconciliation_df.columns, 1):
            val = row[col_name]
            cell = ws_rec.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border_thin
            if col_name == "Status":
                cell.alignment = align_center
                if val == "MATCH":
                    cell.fill = accent_green
                    cell.font = Font(name="Calibri", size=10, bold=True, color="047857")
                else:
                    cell.fill = accent_red
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    auto_fit_columns(ws_rec)

    # =========================================================================
    # 10. LLM VALIDATION SHEET
    # =========================================================================
    ws_llm = wb.create_sheet(title="LLM Validation")
    ws_llm.views.sheetView[0].showGridLines = True

    ws_llm.merge_cells("A1:D1")
    ws_llm["A1"] = f"LLM INDEPENDENT VALIDATION AUDIT ({llm_audit.get('llm_engine', 'QuantLab Engine')})"
    ws_llm["A1"].fill = primary_fill
    ws_llm["A1"].font = font_title
    ws_llm["A1"].alignment = align_center

    lines = llm_text.splitlines()
    for r_idx, line in enumerate(lines, 3):
        cell = ws_llm.cell(row=r_idx, column=1, value=line)
        cell.font = font_code

    ws_llm.column_dimensions["A"].width = 110

    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
